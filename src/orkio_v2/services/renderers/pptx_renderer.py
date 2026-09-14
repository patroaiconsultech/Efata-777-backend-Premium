from __future__ import annotations
import io
from ..artifact_specs import ArtifactDocumentSpec,build_slide_plan

class PptxRendererUnavailable(RuntimeError): pass

def render_pptx(spec:ArtifactDocumentSpec)->bytes:
    try:
        from pptx import Presentation
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches,Pt
    except ImportError as exc:
        raise PptxRendererUnavailable("ARTIFACT_PPTX_RENDERER_UNAVAILABLE") from exc
    deck_title,slides=build_slide_plan(spec)
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    prs.core_properties.title=deck_title; prs.core_properties.author="EFATÀ 777"
    title_slide=prs.slides.add_slide(prs.slide_layouts[0]); title_slide.shapes.title.text=deck_title
    title_slide.placeholders[1].text="EFATÀ 777"
    title_slide.shapes.title.text_frame.paragraphs[0].font.size=Pt(30)
    for item in slides:
        slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text=item.title
        slide.shapes.title.text_frame.paragraphs[0].font.size=Pt(25)
        frame=slide.placeholders[1].text_frame; frame.clear(); frame.word_wrap=True
        size=20 if len(item.bullets)<=4 else 17
        for idx,bullet in enumerate(item.bullets):
            p=frame.paragraphs[0] if idx==0 else frame.add_paragraph()
            p.text=bullet; p.level=0; p.font.size=Pt(size); p.space_after=Pt(8)
        footer=slide.shapes.add_textbox(Inches(10.9),Inches(7.02),Inches(2),Inches(.25))
        footer.text_frame.text="EFATÀ 777"; footer.text_frame.paragraphs[0].alignment=PP_ALIGN.RIGHT; footer.text_frame.paragraphs[0].font.size=Pt(8)
    out=io.BytesIO(); prs.save(out); return out.getvalue()
