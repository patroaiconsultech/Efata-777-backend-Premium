from __future__ import annotations
import io
from ..artifact_specs import ArtifactDocumentSpec,spreadsheet_rows

class XlsxRendererUnavailable(RuntimeError): pass

def render_xlsx(content:str,spec:ArtifactDocumentSpec)->bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment,Font,PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise XlsxRendererUnavailable("ARTIFACT_XLSX_RENDERER_UNAVAILABLE") from exc
    rows=spreadsheet_rows(content,spec)
    if not rows: raise ValueError("ARTIFACT_EMPTY_CONTENT")
    wb=Workbook(); ws=wb.active; ws.title="EFATÀ 777"
    for row in rows: ws.append(row)
    max_cols=max(len(r) for r in rows)
    if len(rows)>1:
        ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(max_cols)}{len(rows)}"
    fill=PatternFill("solid",fgColor="E9EEF5")
    for cell in ws[1]:
        cell.font=Font(bold=True); cell.fill=fill; cell.alignment=Alignment(vertical="top",wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
    for idx in range(1,max_cols+1):
        vals=[str(ws.cell(r,idx).value or "") for r in range(1,len(rows)+1)]
        ws.column_dimensions[get_column_letter(idx)].width=min(max(max((len(v) for v in vals),default=8)+2,10),42)
    ws.sheet_view.showGridLines=False
    out=io.BytesIO(); wb.save(out); return out.getvalue()
