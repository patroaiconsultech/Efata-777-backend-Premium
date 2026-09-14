from __future__ import annotations

import io
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from conftest import headers
from orkio_v2.config import get_settings
from orkio_v2.services.artifact_generation import (
    artifact_capability_manifest,
    artifact_generation_system_message,
    detect_artifact_intent,
    render_and_validate,
)
from orkio_v2.services.artifact_specs import (
    HeadingBlock,
    ListBlock,
    TableBlock,
    parse_document_spec,
)
from orkio_v2.services.image_providers.base import ImageGenerationResult
from orkio_v2.services.renderers.image_renderer import validate_image_result


SAMPLE = """# Relatório EFATÀ 777

Resumo executivo com evidência objetiva.

## Capacidades

- Realtime governado
- Document Intake
- Artifact Generation

| Capacidade | Estado | Evidência |
|---|---|---|
| DOCX | Premium | reopen |
| PDF | Premium | reopen |
| PPTX | Premium | multi-slide |
| XLSX | Premium | workbook |

[[PAGE_BREAK]]

## Próximos passos

1. Auditar
2. Validar
3. Congelar
"""


def test_r4_structured_spec_recognizes_headings_lists_table_and_break():
    spec = parse_document_spec(SAMPLE)
    assert spec.title == "Relatório EFATÀ 777"
    assert any(isinstance(block, HeadingBlock) for block in spec.blocks)
    assert any(isinstance(block, ListBlock) for block in spec.blocks)
    table = next(block for block in spec.blocks if isinstance(block, TableBlock))
    assert table.rows[0] == ("Capacidade", "Estado", "Evidência")
    assert "Artifact Generation" in spec.semantic_text


def test_r4_docx_premium_has_styles_numbering_table_header_footer_and_reopens():
    intent = detect_artifact_intent("gere docx")
    out = render_and_validate(intent=intent, content=SAMPLE, filename="probe.docx")
    assert out.renderer == "efata_docx_premium_v2"
    with zipfile.ZipFile(io.BytesIO(out.data)) as z:
        assert z.testzip() is None
        names = set(z.namelist())
        assert {"word/styles.xml", "word/numbering.xml", "word/header1.xml", "word/footer1.xml"} <= names
        document = z.read("word/document.xml").decode("utf-8")
        assert 'w:val="Heading1"' in document
        assert "<w:tbl>" in document
        assert 'w:type="page"' in document
        assert "Relatório EFATÀ 777" in document
        assert "EFATÀ 777" in z.read("word/header1.xml").decode("utf-8")


def test_r4_pdf_flowable_wraps_and_paginates_long_content():
    intent = detect_artifact_intent("gere PDF")
    long_content = "# Relatório\n\n" + "\n\n".join(
        f"## Seção {i}\n" + ("Conteúdo técnico governado com evidência e rastreabilidade. " * 20)
        for i in range(1, 9)
    )
    out = render_and_validate(intent=intent, content=long_content, filename="probe.pdf")
    assert out.renderer == "efata_pdf_flowable_v2"
    reader = PdfReader(io.BytesIO(out.data), strict=True)
    assert len(reader.pages) >= 2
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Relatório" in text and "Seção 8" in text


def test_r4_pptx_is_multislide_and_bounded():
    intent = detect_artifact_intent("gere pptx")
    out = render_and_validate(intent=intent, content=SAMPLE, filename="probe.pptx")
    assert out.renderer == "efata_pptx_multislide_v2"
    prs = Presentation(io.BytesIO(out.data))
    assert len(prs.slides) >= 3
    content_slides = list(prs.slides)[1:]
    for slide in content_slides:
        body_shapes = [s for s in slide.shapes if getattr(s, "has_text_frame", False)]
        joined = "\n".join(s.text for s in body_shapes)
        assert len(joined) < 1800


def test_r4_xlsx_has_values_style_freeze_and_filter():
    intent = detect_artifact_intent("gere xlsx")
    out = render_and_validate(intent=intent, content=SAMPLE, filename="probe.xlsx")
    assert out.renderer == "efata_xlsx_premium_v2"
    wb = load_workbook(io.BytesIO(out.data))
    ws = wb.active
    assert ws["A1"].value == "Capacidade"
    assert ws["B2"].value == "Premium"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert ws["A1"].font.bold is True
    wb.close()


def test_r4_image_boundary_validates_bytes_without_real_provider():
    result = ImageGenerationResult(
        data=b"\x89PNG\r\n\x1a\n" + b"x" * 32,
        mime_type="image/png",
        provider="test-only",
        model="test-only",
        width=64,
        height=64,
    )
    ext, digest = validate_image_result(result, max_bytes=1024)
    assert ext == ".png"
    assert len(digest) == 64


def test_r4_artifact_capabilities_are_conservative(client, monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "artifacts_enabled", True, raising=False)
    response = client.get("/api/v2/artifacts/capabilities", headers=headers())
    assert response.status_code == 200
    body = response.json()
    assert body["formats"]["docx"]["supported"] is True
    assert body["formats"]["pptx"]["enabled"] is True
    assert body["formats"]["image"]["supported"] is False
    assert body["formats"]["image"]["provider_configured"] is False
    assert body["runtime_proved"] is False
    assert body["premium_qualified"] is False


def test_r4_format_specific_system_prompt_preserves_governance_language():
    intent = detect_artifact_intent("gere uma apresentação pptx")
    msg = artifact_generation_system_message(intent)["content"]
    assert "Produce only the complete document body" in msg
    assert "## for slide sections" in msg
    assert "Do not invent a download URL" in msg
